"""
Tests for the instant, single-asset actions in asset_actions.py:
table -> rows conversion (no LLM, no whole-paper job) and CSV/XLSX/JSON
export, available as soon as one table has been extracted rather than
after the whole paper finishes.
"""
import io
import shutil
from pathlib import Path

import openpyxl
import pytest

from app.db.models import ExtractionAsset, Paper, Project
from app.core.security import create_access_token, hash_password
from app.db.models import User

FIXTURE_CSV = Path(__file__).parent / "fixtures" / "sample_table.csv"


def _seed_table_asset(db, user, tmp_path, caption="Yeasts and molds counts in Gouda cheese stored at 4°C"):
    project = Project(name="Cheese Shelf-Life", owner_id=user.id)
    db.add(project)
    db.flush()

    paper = Paper(project_id=project.id, filename="p.pdf", original_name="paper.pdf",
                  file_path="/nonexistent/p.pdf", status="extracted")
    db.add(paper)
    db.flush()

    csv_path = tmp_path / "table_0.csv"
    shutil.copy(FIXTURE_CSV, csv_path)

    asset = ExtractionAsset(
        paper_id=paper.id, project_id=project.id, docling_item_ref="#/tables/0",
        asset_type="native_table", page_number=5, classification="native_table",
        caption=caption, csv_path=str(csv_path), csv_rows=2, csv_cols=4,
        relevance_score=5.0,
    )
    db.add(asset)
    db.commit()
    db.refresh(project)
    db.refresh(paper)
    db.refresh(asset)
    return project, paper, asset


def _seed_figure_asset(db, user):
    project = Project(name="Cheese Shelf-Life", owner_id=user.id)
    db.add(project)
    db.flush()
    paper = Paper(project_id=project.id, filename="p.pdf", original_name="paper.pdf",
                  file_path="/nonexistent/p.pdf", status="extracted")
    db.add(paper)
    db.flush()
    asset = ExtractionAsset(
        paper_id=paper.id, project_id=project.id, docling_item_ref="#/pictures/0",
        asset_type="figure", page_number=2, classification="chart",
    )
    db.add(asset)
    db.commit()
    db.refresh(project)
    db.refresh(paper)
    db.refresh(asset)
    return project, paper, asset


def test_extract_rows_time_series_table(client, auth_headers, db, test_user, tmp_path):
    project, paper, asset = _seed_table_asset(db, test_user, tmp_path)

    resp = client.post(
        f"/api/projects/{project.id}/papers/{paper.id}/assets/{asset.id}/extract-rows",
        headers=auth_headers,
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["experiment_count"] == 2      # Control + EO 0.5%
    assert body["row_count"] == 6              # 2 rows x 3 time points
    assert {r["treatment"].lower() for r in body["rows"]} >= {"control"}
    for row in body["rows"]:
        assert row["indicator_type"] == "yeasts and molds"
        assert row["confidence"] is not None


def test_extract_rows_rejects_figure_asset(client, auth_headers, db, test_user):
    project, paper, asset = _seed_figure_asset(db, test_user)

    resp = client.post(
        f"/api/projects/{project.id}/papers/{paper.id}/assets/{asset.id}/extract-rows",
        headers=auth_headers,
    )
    assert resp.status_code == 400


def test_extract_rows_requires_project_access(client, db, test_user, tmp_path):
    project, paper, asset = _seed_table_asset(db, test_user, tmp_path)

    other = User(email="asset-other@example.com", full_name="Other",
                 hashed_password=hash_password("password123"), is_active=True)
    db.add(other)
    db.commit()
    db.refresh(other)
    other_headers = {"Authorization": f"Bearer {create_access_token({'sub': str(other.id)})}"}

    resp = client.post(
        f"/api/projects/{project.id}/papers/{paper.id}/assets/{asset.id}/extract-rows",
        headers=other_headers,
    )
    assert resp.status_code == 404


def test_export_csv_returns_the_raw_table(client, auth_headers, db, test_user, tmp_path):
    project, paper, asset = _seed_table_asset(db, test_user, tmp_path)

    resp = client.get(
        f"/api/projects/{project.id}/papers/{paper.id}/assets/{asset.id}/export?format=csv",
        headers=auth_headers,
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("text/csv")
    assert "Control" in resp.text


def test_export_json_includes_table_and_rows(client, auth_headers, db, test_user, tmp_path):
    project, paper, asset = _seed_table_asset(db, test_user, tmp_path)

    resp = client.get(
        f"/api/projects/{project.id}/papers/{paper.id}/assets/{asset.id}/export?format=json",
        headers=auth_headers,
    )
    assert resp.status_code == 200
    body = resp.json()
    assert len(body["table"]) == 2
    assert len(body["rows"]) == 6


def test_export_xlsx_returns_a_two_sheet_workbook(client, auth_headers, db, test_user, tmp_path):
    project, paper, asset = _seed_table_asset(db, test_user, tmp_path)

    resp = client.get(
        f"/api/projects/{project.id}/papers/{paper.id}/assets/{asset.id}/export?format=xlsx",
        headers=auth_headers,
    )
    assert resp.status_code == 200
    assert "spreadsheetml" in resp.headers["content-type"]
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["Table", "Rows"]
    assert wb["Table"].max_row == 3   # header + 2 rows
    assert wb["Rows"].max_row == 7    # header + 6 rows


def test_export_json_handles_blank_cells_without_500(client, auth_headers, db, test_user, tmp_path):
    """Regression test for a real bug found live in the UI: a table with a
    blank cell (very common — e.g. a day-42 reading missing for one sample)
    reads as NaN via pandas, and NaN is not valid JSON — FastAPI's encoder
    raised 'ValueError: Out of range float values are not JSON compliant:
    nan', 500ing the whole export and breaking the table preview in
    AssetDetailPanel. Blank cells must become null in the response, not
    crash the endpoint."""
    project = Project(name="Cheese Shelf-Life", owner_id=test_user.id)
    db.add(project)
    db.flush()
    paper = Paper(project_id=project.id, filename="p.pdf", original_name="paper.pdf",
                  file_path="/nonexistent/p.pdf", status="extracted")
    db.add(paper)
    db.flush()

    csv_path = tmp_path / "table_blank_cell.csv"
    csv_path.write_text("Treatment,Day 1,Day 42\nControl,3.9,3.7\nEXg,4.2,\n", encoding="utf-8")

    asset = ExtractionAsset(
        paper_id=paper.id, project_id=project.id, docling_item_ref="#/tables/1",
        asset_type="native_table", page_number=7, classification="native_table",
        caption="Table with a missing day-42 reading", csv_path=str(csv_path),
        csv_rows=2, csv_cols=3, relevance_score=5.0,
    )
    db.add(asset)
    db.commit()
    db.refresh(asset)

    resp = client.get(
        f"/api/projects/{project.id}/papers/{paper.id}/assets/{asset.id}/export?format=json",
        headers=auth_headers,
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert len(body["table"]) == 2
    blank_row = next(r for r in body["table"] if r["Treatment"] == "EXg")
    assert blank_row["Day 42"] is None


def test_export_requires_project_access(client, db, test_user, tmp_path):
    project, paper, asset = _seed_table_asset(db, test_user, tmp_path)

    other = User(email="export-other@example.com", full_name="Other",
                 hashed_password=hash_password("password123"), is_active=True)
    db.add(other)
    db.commit()
    db.refresh(other)
    other_headers = {"Authorization": f"Bearer {create_access_token({'sub': str(other.id)})}"}

    resp = client.get(
        f"/api/projects/{project.id}/papers/{paper.id}/assets/{asset.id}/export?format=csv",
        headers=other_headers,
    )
    assert resp.status_code == 404
